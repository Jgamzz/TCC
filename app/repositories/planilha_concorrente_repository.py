import re
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app import db
from app.config import Config
from app.models.preco_concorrente_planilha import PrecoConcorrentePlanilha


class PlanilhaConcorrenteRepository:
    COLUNAS_PRODUTO = {
        "produto",
        "nomeproduto",
        "nomedoproduto",
        "nome",
        "descricao",
        "item",
        "name",
        "product",
        "productname",
    }
    COLUNAS_VALOR = {
        "valor",
        "preco",
        "valorproduto",
        "valordoproduto",
        "value",
        "precoproduto",
        "precodoproduto",
        "price",
        "unitprice",
        "valorunitario",
    }

    def __init__(self, pasta=None):
        self.pasta = Path(pasta or Config.PLANILHAS_CONCORRENTES_DIR)

    @staticmethod
    def _normalizar_cabecalho(valor):
        texto = unicodedata.normalize("NFKD", str(valor or ""))
        texto = "".join(
            caractere for caractere in texto if not unicodedata.combining(caractere)
        )
        return re.sub(r"[^a-z0-9]", "", texto.lower())

    @staticmethod
    def _converter_valor(valor, arquivo, aba, linha):
        if isinstance(valor, bool) or valor is None:
            raise ValueError(
                f"Valor inválido em {arquivo}, aba '{aba}', linha {linha}."
            )
        if isinstance(valor, (int, float, Decimal)):
            try:
                convertido = Decimal(str(valor))
            except InvalidOperation as error:
                raise ValueError(
                    f"Valor inválido em {arquivo}, aba '{aba}', linha {linha}."
                ) from error
        else:
            texto = str(valor).strip().replace("R$", "").replace(" ", "")
            if "," in texto:
                texto = texto.replace(".", "").replace(",", ".")
            try:
                convertido = Decimal(texto)
            except InvalidOperation as error:
                raise ValueError(
                    f"Valor inválido em {arquivo}, aba '{aba}', linha {linha}."
                ) from error
        if not convertido.is_finite() or convertido < 0:
            raise ValueError(
                f"Valor inválido em {arquivo}, aba '{aba}', linha {linha}."
            )
        return convertido.quantize(Decimal("0.01"))

    def _ler_arquivo(self, caminho):
        workbook = load_workbook(caminho, data_only=True, read_only=True)
        registros = []
        try:
            for sheet in workbook.worksheets:
                linhas = enumerate(sheet.iter_rows(values_only=True), start=1)
                cabecalho = None
                coluna_produto = None
                coluna_valor = None
                linha_cabecalho = None

                for numero_linha, valores in linhas:
                    colunas = {
                        self._normalizar_cabecalho(valor): indice
                        for indice, valor in enumerate(valores)
                    }
                    coluna_produto = next(
                        (
                            indice
                            for nome, indice in colunas.items()
                            if nome in self.COLUNAS_PRODUTO
                        ),
                        None,
                    )
                    coluna_valor = next(
                        (
                            indice
                            for nome, indice in colunas.items()
                            if nome in self.COLUNAS_VALOR
                        ),
                        None,
                    )
                    if coluna_produto is not None and coluna_valor is not None:
                        cabecalho = valores
                        linha_cabecalho = numero_linha
                        break

                if cabecalho is None:
                    if sheet.max_row == 1 and sheet.max_column == 1:
                        continue
                    raise ValueError(
                        f"A planilha '{caminho.name}', aba '{sheet.title}', "
                        "precisa ter uma linha de cabeçalho para produto e valor "
                        "(por exemplo: Nome e Valor)."
                    )

                if linha_cabecalho is None:
                    continue

                for numero_linha, linha in enumerate(
                    (valores for _, valores in linhas),
                    start=linha_cabecalho + 1,
                ):
                    if max(coluna_produto, coluna_valor) >= len(linha):
                        continue
                    nome_produto = linha[coluna_produto]
                    valor = linha[coluna_valor]
                    if nome_produto is None and valor is None:
                        continue
                    if nome_produto is None or not str(nome_produto).strip():
                        raise ValueError(
                            f"Produto ausente em {caminho.name}, aba "
                            f"'{sheet.title}', linha {numero_linha}."
                        )

                    registros.append(
                        {
                            "nome_arquivo": caminho.name,
                            "nome_aba": sheet.title,
                            "linha_excel": numero_linha,
                            "nome_produto": str(nome_produto).strip(),
                            "valor": self._converter_valor(
                                valor, caminho.name, sheet.title, numero_linha
                            ),
                        }
                    )
        finally:
            workbook.close()
        return registros

    def sincronizar(self):
        self.pasta.mkdir(parents=True, exist_ok=True)
        arquivos = sorted(
            caminho
            for caminho in self.pasta.iterdir()
            if caminho.is_file()
            and caminho.suffix.lower() in {".xlsx", ".xlsm"}
            and not caminho.name.startswith("~$")
        )
        registros = [
            registro
            for arquivo in arquivos
            for registro in self._ler_arquivo(arquivo)
        ]

        arquivos_lidos = {arquivo.name for arquivo in arquivos}
        registros_existentes = db.session.scalars(
            select(PrecoConcorrentePlanilha)
        ).all()
        existentes_por_origem = {
            (registro.nome_arquivo, registro.nome_aba, registro.linha_excel): registro
            for registro in registros_existentes
        }

        origens_lidas = set()
        for chave, registro in existentes_por_origem.items():
            if registro.nome_arquivo not in arquivos_lidos:
                db.session.delete(registro)
                continue

        for dados in registros:
            chave = (
                dados["nome_arquivo"],
                dados["nome_aba"],
                dados["linha_excel"],
            )
            origens_lidas.add(chave)
            registro = existentes_por_origem.get(chave)
            if registro is None:
                db.session.add(PrecoConcorrentePlanilha(**dados))
            else:
                registro.nome_produto = dados["nome_produto"]
                registro.valor = dados["valor"]

        for chave, registro in existentes_por_origem.items():
            if chave not in origens_lidas:
                db.session.delete(registro)

        db.session.commit()
        return len(arquivos), len(registros)

    def adicionar_produto(self, nome_produto, precos):
        """Adiciona ou atualiza um produto nas planilhas dos concorrentes."""
        nome_produto = str(nome_produto or "").strip()
        if not nome_produto:
            raise ValueError("Informe o nome do produto.")
        if not isinstance(precos, dict) or not precos:
            raise ValueError("Informe ao menos um preço de concorrente.")

        arquivos_atualizados = 0
        for nome_arquivo, valor in precos.items():
            if valor in (None, ""):
                continue
            caminho = self.pasta / str(nome_arquivo)
            if (
                caminho.parent != self.pasta
                or not caminho.is_file()
                or caminho.suffix.lower() not in {".xlsx", ".xlsm"}
            ):
                raise ValueError(f"Planilha de concorrente inválida: {nome_arquivo}.")
            preco = self._converter_valor(valor, caminho.name, "", 0)
            workbook = load_workbook(
                caminho,
                keep_vba=caminho.suffix.lower() == ".xlsm",
            )
            try:
                sheet = workbook.worksheets[0]
                coluna_produto = None
                coluna_valor = None
                linha_cabecalho = None
                for numero_linha, linha in enumerate(
                    sheet.iter_rows(values_only=True), start=1
                ):
                    colunas = {
                        self._normalizar_cabecalho(valor): indice + 1
                        for indice, valor in enumerate(linha)
                    }
                    coluna_produto = next(
                        (
                            indice
                            for nome, indice in colunas.items()
                            if nome in self.COLUNAS_PRODUTO
                        ),
                        None,
                    )
                    coluna_valor = next(
                        (
                            indice
                            for nome, indice in colunas.items()
                            if nome in self.COLUNAS_VALOR
                        ),
                        None,
                    )
                    if coluna_produto and coluna_valor:
                        linha_cabecalho = numero_linha
                        break
                if linha_cabecalho is None:
                    raise ValueError(
                        f"A planilha '{caminho.name}' não possui cabeçalho de produto e valor."
                    )

                linha_produto = None
                chave_produto = self._normalizar_cabecalho(nome_produto)
                for numero_linha in range(linha_cabecalho + 1, sheet.max_row + 1):
                    valor_produto = sheet.cell(numero_linha, coluna_produto).value
                    if self._normalizar_cabecalho(valor_produto) == chave_produto:
                        linha_produto = numero_linha
                        break
                if linha_produto is None:
                    linha_produto = sheet.max_row + 1
                    sheet.cell(linha_produto, coluna_produto).value = nome_produto
                sheet.cell(linha_produto, coluna_valor).value = float(preco)
                workbook.save(caminho)
                arquivos_atualizados += 1
            finally:
                workbook.close()

        if not arquivos_atualizados:
            raise ValueError("Informe ao menos um preço válido.")
        self.sincronizar()
        return arquivos_atualizados

    def listar_comparacao(self):
        registros = db.session.scalars(
            select(PrecoConcorrentePlanilha).order_by(
                PrecoConcorrentePlanilha.nome_arquivo,
                PrecoConcorrentePlanilha.nome_produto,
                PrecoConcorrentePlanilha.linha_excel,
            )
        ).all()

        arquivos = sorted({registro.nome_arquivo for registro in registros})
        produtos = {}
        chaves_produtos = {}
        for registro in registros:
            normalizado = self._normalizar_cabecalho(registro.nome_produto)
            if normalizado not in produtos:
                produtos[normalizado] = registro.nome_produto
                chaves_produtos[normalizado] = {}
            chaves_produtos[normalizado].setdefault(
                registro.nome_arquivo, []
            ).append(float(registro.valor))

        comparacoes = []
        for chave, nome_produto in produtos.items():
            precos = {
                arquivo: min(valores)
                for arquivo, valores in chaves_produtos[chave].items()
            }
            valores = list(precos.values())
            comparacoes.append(
                {
                    "nome_produto": nome_produto,
                    "precos": {
                        arquivo: precos.get(arquivo) for arquivo in arquivos
                    },
                    "diferenca": (
                        round(max(valores) - min(valores), 2)
                        if len(valores) >= 2
                        else None
                    ),
                }
            )

        comparacoes.sort(key=lambda produto: produto["nome_produto"].casefold())
        return {"arquivos": arquivos, "produtos": comparacoes}
