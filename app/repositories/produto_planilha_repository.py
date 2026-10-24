from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook, load_workbook
from sqlalchemy import or_, select

from app import db
from app.config import Config
from app.models.estoque import Estoque
from app.models.produto import Produto


class ProdutoPlanilhaRepository:
    CABECALHOS = (
        "codigo",
        "nome",
        "descricao",
        "categoria",
        "fornecedor",
        "custo",
        "preco_atual",
        "estoque",
        "data_validade",
        "ativo",
    )

    def __init__(self, caminho=None):
        self.caminho = Path(caminho or Config.PRODUTOS_FILE)

    @staticmethod
    def _decimal(valor, campo):
        try:
            convertido = Decimal(str(valor))
        except (InvalidOperation, TypeError, ValueError):
            raise ValueError(f"{campo} deve ser numérico na planilha de produtos.") from None
        if not convertido.is_finite():
            raise ValueError(f"{campo} deve ser finito na planilha de produtos.")
        return convertido

    @staticmethod
    def _data(valor):
        if valor in (None, ""):
            return None
        if isinstance(valor, datetime):
            return valor.date()
        if isinstance(valor, date):
            return valor
        try:
            return datetime.strptime(str(valor).strip(), "%Y-%m-%d").date()
        except ValueError:
            raise ValueError(
                "data_validade deve estar no formato AAAA-MM-DD na planilha de produtos."
            ) from None

    @staticmethod
    def _booleano(valor):
        if valor in (None, ""):
            return True
        return str(valor).strip().lower() not in {"0", "false", "nao", "não", "inativo"}

    def _exportar_atual(self):
        self.caminho.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Produtos"
        sheet.append(self.CABECALHOS)
        for produto in db.session.scalars(select(Produto).order_by(Produto.id_produto)):
            sheet.append(
                [
                    produto.codigo or f"LEGADO-{produto.id_produto}",
                    produto.nome,
                    produto.descricao,
                    produto.categoria,
                    produto.fornecedor,
                    float(produto.custo),
                    float(produto.preco_atual),
                    float(produto.estoque),
                    produto.data_validade,
                    produto.ativo,
                ]
            )
        workbook.save(self.caminho)
        workbook.close()

    def sincronizar(self):
        if not self.caminho.is_file():
            self._exportar_atual()
            return 0

        workbook = load_workbook(self.caminho, data_only=True, read_only=True)
        try:
            sheet = workbook.active
            linhas = sheet.iter_rows(values_only=True)
            cabecalho = next(linhas, None)
            if not cabecalho:
                return 0
            colunas = {
                str(valor or "").strip().lower(): indice
                for indice, valor in enumerate(cabecalho)
            }
            ausentes = {"nome", "custo", "preco_atual", "estoque"} - set(colunas)
            if ausentes:
                raise ValueError(
                    "A planilha de produtos não possui as colunas: "
                    + ", ".join(sorted(ausentes))
                )

            importados = 0
            for valores in linhas:
                dados = {
                    campo: valores[indice] if indice < len(valores) else None
                    for campo, indice in colunas.items()
                }
                nome = str(dados.get("nome") or "").strip()
                if not nome:
                    continue
                codigo = str(dados.get("codigo") or "").strip().upper() or None
                existente = db.session.scalar(
                    select(Produto).where(
                        or_(
                            Produto.nome.ilike(nome),
                            Produto.codigo == codigo if codigo else False,
                        )
                    )
                )
                if existente:
                    continue

                estoque = self._decimal(dados.get("estoque", 0), "estoque")
                produto = Produto(
                    codigo=codigo,
                    nome=nome,
                    descricao=str(dados.get("descricao") or "").strip() or None,
                    categoria=str(dados.get("categoria") or "").strip() or None,
                    fornecedor=str(dados.get("fornecedor") or "").strip() or None,
                    custo=self._decimal(dados.get("custo"), "custo"),
                    preco_atual=self._decimal(dados.get("preco_atual"), "preco_atual"),
                    estoque=estoque,
                    data_validade=self._data(dados.get("data_validade")),
                    ativo=self._booleano(dados.get("ativo")),
                )
                db.session.add(produto)
                db.session.flush()
                if estoque > 0:
                    db.session.add(
                        Estoque(
                            id_produto=produto.id_produto,
                            tipo_operacao="ENTRADA",
                            quantidade=estoque,
                            motivo="Estoque inicial importado da planilha de produtos",
                        )
                    )
                importados += 1
            db.session.commit()
            return importados
        except Exception:
            db.session.rollback()
            raise
        finally:
            workbook.close()